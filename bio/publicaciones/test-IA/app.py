"""TEST-IA: interfaz de versiones y trazabilidad, sin API conectada."""
import io
import streamlit as st
from docx import Document
from instrucciones import INSTRUCCIONES
from versiones import crear_version_original, crear_version
from iteraciones import crear_solicitud, registrar_carga_offline
from exportacion import crear_zip_auditoria
from flujo import solicitar_iteracion, comparar_offline, validar_evaluacion_final
from cierre import cerrar_autoria
from tablero import resumir_atribuciones
from finalizacion import solicitar_cierre, cierre_evaluado, reabrir_revision

MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
st.set_page_config(page_title='TEST-IA', page_icon='🧪', layout='wide')
st.title('TEST-IA')
st.caption('Entorno de uso de IA y auditoría')
st.markdown("""<style>
.block-container {max-width: 1180px; padding-top: 2rem;}
h1 {letter-spacing: -.035em;}
.testia-section {padding: .85rem 1.1rem; margin: 1.6rem 0 .9rem; border-radius: 12px; border-left: 7px solid var(--accent); background: var(--bg);}
.testia-section h2 {font-size: 1.35rem; margin: 0; color: var(--text-color);}
.testia-section p {font-size: .88rem; margin: .3rem 0 0; color: var(--text-color); opacity: .8;}
</style>""", unsafe_allow_html=True)

def seccion(numero, titulo, descripcion, color, fondo):
    st.markdown(f'<div class="testia-section" style="--accent:{color};--bg:{fondo}"><h2>{numero}. {titulo}</h2><p>{descripcion}</p></div>', unsafe_allow_html=True)


seccion(1, 'Instrucciones', 'Definí la intervención y el alcance sobre el manuscrito.', '#3566ad', 'rgba(53,102,173,.09)')
archivo = st.file_uploader('Manuscrito original (.docx)', type=['docx'], key='original_upload')
if archivo is not None:
    contenido = archivo.getvalue()
    if st.session_state.get('original_firma') != (archivo.name, contenido):
        try:
            original_nuevo = crear_version_original(archivo.name, contenido)
        except Exception as exc:
            st.error(f'No se pudo leer el documento: {exc}')
            st.stop()
        st.session_state['original'] = original_nuevo
        st.session_state['versiones'] = []
        st.session_state['eventos'] = []
        st.session_state['evaluacion'] = None
        st.session_state['cierre'] = None
        st.session_state['original_firma'] = (archivo.name, contenido)

cierre = st.session_state.get('cierre')
original = st.session_state.get('original')
versiones = st.session_state.get('versiones', [])
eventos = st.session_state.get('eventos', [])
ultima = versiones[-1] if versiones else original

if original:
    st.success(f'Manuscrito: {original["nombre"]} · Versión vigente: {ultima["numero"]}')
    with st.expander('Vista preliminar del documento'):
        doc = Document(io.BytesIO(ultima['contenido']))
        st.write(f'{len(doc.paragraphs)} párrafos · {len(doc.tables)} tablas')
        st.text_area('Texto de párrafos', '\n\n'.join(p.text for p in doc.paragraphs), height=220)

categoria = st.selectbox('Instrucción', list(INSTRUCCIONES))
st.write(INSTRUCCIONES[categoria]['descripcion'])
with st.expander('Prompt general que recibirá Astra'):
    st.write(INSTRUCCIONES[categoria]['prompt'])
especifica = st.text_area('Prompt específico del AUTOR', key='especifica', height=110)
alcance = st.radio('Alcance', ['Manuscrito completo', 'Fragmento'], horizontal=True)
ubicacion = st.text_input('Ubicación', key='ubicacion')
if st.button('Llamar Astra', type='primary'):
    if not ultima:
        st.warning('Primero cargá el manuscrito.')
    else:
        try:
            solicitud = crear_solicitud(ultima, categoria,
                INSTRUCCIONES[categoria]['prompt'], especifica, alcance, ubicacion)
            st.session_state['eventos'].append(solicitud)
            st.session_state['cierre'] = None
            st.session_state['evaluacion'] = None
            st.success('Instrucción registrada para la versión vigente.')
        except ValueError as exc:
            st.error(str(exc))

st.divider()
seccion(2, 'Versiones', 'Consultá y descargá cada estado del documento.', '#16836b', 'rgba(22,131,107,.09)')
if not original:
    st.info('Cargá un manuscrito para consultar versiones.')
else:
    st.download_button('Descargar original', original['contenido'],
                       file_name=original['nombre'], mime=MIME)
    for version in versiones:
        st.download_button(f'Versión {version["numero"]} · {version["tipo"]}',
                           version['contenido'], file_name=version['nombre'],
                           mime=MIME, key=f'v_{version["numero"]}')
    st.caption('Las intervenciones editoriales de Astra se identifican mediante anotaciones de sugerencia.')
    st.caption('En traducciones completas no se resaltará todo el manuscrito. Las evaluaciones de coherencia integral pueden incorporarse como observaciones al final.')

st.divider()
seccion(3, 'Iteraciones', 'Continuá la revisión desde la versión vigente o incorporá una edición offline.', '#8d59b6', 'rgba(141,89,182,.09)')
if not original:
    st.info('Primero cargá el manuscrito original.')
else:
    if cierre:
        st.info('Revisión finalizada. Podés reabrirla para continuar con nuevas iteraciones.')
        if st.button('Reabrir revisión'):
            st.session_state['cierre'] = reabrir_revision(cierre)
            st.session_state['evaluacion'] = None
            st.rerun()
    if not cierre:
        st.markdown('**Edición offline y nuevas instrucciones**')
    subida = st.file_uploader('Subir Word editado offline (.docx)', type=['docx'],
                              key=f'offline_{ultima["numero"]}')
    if subida is not None and st.button('Incorporar como nueva versión vigente', disabled=bool(cierre)):
        try:
            nuevo = crear_version(ultima, subida.getvalue(), 'autor_offline')
            if nuevo['sha256'] == ultima['sha256']:
                st.warning('El archivo es idéntico a la versión vigente; no se creó otra versión.')
            else:
                st.session_state['versiones'].append(nuevo)
                evento = registrar_carga_offline(ultima, nuevo)
                evento['comparacion_parrafos'] = comparar_offline(ultima['contenido'], nuevo['contenido'])
                st.session_state['eventos'].append(evento)
                st.session_state['evaluacion'] = None
                st.rerun()
        except Exception as exc:
            st.error(f'No se incorporó el documento: {exc}')
    st.markdown('**Solicitar nuevos cambios a Astra**')
    categoria_iter = st.selectbox('Instrucción de la nueva iteración', list(INSTRUCCIONES), key='categoria_iter')
    nuevas_instrucciones = st.text_area('Nuevos cambios solicitados por el AUTOR',
                                        height=180, key='nuevas_instrucciones',
                                        placeholder='Ej.: Reforzar el argumento del apartado 3 y revisar las citas.')
    alcance_iter = st.radio('Alcance de la nueva iteración',
                           ['Manuscrito completo', 'Fragmento'], horizontal=True, key='alcance_iter')
    ubicacion_iter = st.text_input('Ubicación del fragmento (si corresponde)', key='ubicacion_iter')
    if st.button('Llamar Astra · nueva iteración', disabled=bool(cierre)):
        try:
            solicitud = solicitar_iteracion(ultima, categoria_iter,
                INSTRUCCIONES[categoria_iter]['prompt'], nuevas_instrucciones,
                alcance_iter, ubicacion_iter)
            st.session_state['eventos'].append(solicitud)
            st.session_state['evaluacion'] = None
            st.success(f'Instrucciones registradas sobre la Versión {ultima["numero"]}.')
        except ValueError as exc:
            st.error(str(exc))
    st.divider()
    if st.button('Finalizar revisión y evaluar atribución', type='primary', disabled=bool(cierre)):
        st.session_state['cierre'] = solicitar_cierre(ultima, st.session_state['eventos'])
        st.session_state['evaluacion'] = None
        st.session_state['eventos'].append({
            'tipo': 'Finalización solicitada',
            'fecha_utc': st.session_state['cierre']['fecha_utc'],
            'base_version': ultima['numero'],
            'base_sha256': ultima['sha256'],
            'estado': 'Evaluación final solicitada',
        })
        st.rerun()
    st.write(f'Eventos registrados: {len(eventos)}')
    with st.expander('Historial interno de iteraciones'):
        for evento in eventos:
            st.write(f'{evento["fecha_utc"]} · Base {evento["base_version"]} · '
                     f'{evento.get("tipo", evento.get("categoria", "Solicitud"))} · '
                     f'{evento.get("estado", "Registrado")}')

st.divider()
seccion(4, 'Tablero', 'Distribución de contribuciones y estimaciones de atribución.', '#bd7728', 'rgba(189,119,40,.09)')
evaluacion = st.session_state.get('evaluacion')
evaluacion_vigente = cierre_evaluado(st.session_state.get('cierre'), ultima, evaluacion)
resumen = resumir_atribuciones(evaluacion if evaluacion_vigente else None)
a, b, c = st.columns(3)
a.metric('Anotaciones de atribución total', resumen['verde'] if evaluacion_vigente else '—')
b.metric('Atribuciones en incertidumbre', resumen['amarillo'] if evaluacion_vigente else '—')
c.metric('Total atribuido', resumen['total'] if evaluacion_vigente else '—')
st.markdown('**Atribuciones en incertidumbre · distribución por cuartiles**')
columnas = st.columns(4)
rotulos = ('Q1 · 0–25 %', 'Q2 · >25–50 %', 'Q3 · >50–75 %', 'Q4 · >75–100 %')
for columna, rotulo, cantidad in zip(columnas, rotulos, resumen['cuartiles']):
    columna.metric(rotulo, cantidad if evaluacion_vigente else '—')
if evaluacion_vigente:
    if resumen['sin_porcentaje']:
        st.caption(f'{resumen["sin_porcentaje"]} contribución(es) sin cuartil asignado.')
    st.caption('Los cuartiles se aplican exclusivamente a atribuciones en incertidumbre; expresan la clasificación por intervalo, sin porcentajes puntuales.')
else:
    st.caption('El Tablero presenta los resultados de atribución del documento final, una vez concluida la revisión.')

st.divider()
seccion(5, 'Descargas', 'Documento anotado y archivo completo de auditoría.', '#4a7894', 'rgba(74,120,148,.09)')
if not original:
    st.info('Cargá un documento para habilitar la auditoría.')
else:
    word_final = None
    if evaluacion_vigente:
        try:
            word_final, _ = cerrar_autoria(ultima, evaluacion)
            st.download_button('Descargar Documento Anotado', word_final,
                               file_name='TEST-IA_final.docx', mime=MIME)
        except ValueError as exc:
            st.error(f'Evaluación final inválida: {exc}')
    if evaluacion_vigente and word_final is not None:
        zip_bytes = crear_zip_auditoria(original, versiones, eventos, evaluacion, word_final)
        st.download_button('Descargar ZIP de auditoría', zip_bytes,
                           file_name='TEST-IA_auditoria.zip', mime='application/zip')
