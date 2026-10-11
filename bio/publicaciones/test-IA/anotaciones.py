"""Generación del documento anotado de TEST-IA."""

from io import BytesIO

from docx import Document


def crear_documento_anotado(original, registros):
    """Genera un DOCX anotado sin modificar el original."""
    documento = Document(BytesIO(original["contenido"]))

    documento.add_page_break()
    documento.add_heading("Anotaciones de IA", level=1)

    if not registros:
        documento.add_paragraph(
            "No hay intervenciones registradas."
        )

    for numero, registro in enumerate(registros, start=1):
        documento.add_heading(
            f'Intervención {numero}: {registro["categoria"]}',
            level=2,
        )

        documento.add_paragraph(
            f'Documento: {registro["documento"]}'
        )
        documento.add_paragraph(
            f'Instrucción: {registro["instruccion_especifica"]}'
        )
        documento.add_paragraph(
            f'Ubicación: {registro["ubicacion"] or "No especificada"}'
        )
        documento.add_paragraph(
            f'Propuesta de IA: '
            f'{registro["propuesta_ia"] or "Pendiente de ejecución"}'
        )
        documento.add_paragraph(
            f'Decisión del autor: {registro["decision_autor"]}'
        )
        documento.add_paragraph(
            f'Observaciones: {registro["observaciones_autor"]}'
        )

    salida = BytesIO()
    documento.save(salida)
    return salida.getvalue()
