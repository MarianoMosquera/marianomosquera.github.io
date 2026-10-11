"""Cierre de autoría: verifica evaluación, documento base y resultado sin API."""
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from flujo import validar_evaluacion_final


def _textos(doc):
    return ([p.text for p in doc.paragraphs],
            [[[p.text for p in celda.paragraphs] for fila in t.rows for celda in fila.cells]
             for t in doc.tables])


def cerrar_autoria(ultima, evaluacion):
    """Exige identificación de la versión evaluada y procedencia explícita de Astra.

    La función valida atribuciones recibidas, pero no las infiere.
    """
    if not isinstance(evaluacion, dict):
        raise ValueError('Falta evaluación final de Astra')
    if evaluacion.get('evaluador') != 'astra':
        raise ValueError('La evaluación final debe proceder de Astra')
    if evaluacion.get('version_evaluada') != ultima['numero'] or evaluacion.get('sha256_evaluado') != ultima['sha256']:
        raise ValueError('La evaluación no corresponde a la última versión vigente')
    if not evaluacion.get('justificacion_global', '').strip():
        raise ValueError('Falta justificación global de la evaluación')
    resultado = validar_evaluacion_final(evaluacion, ultima['contenido'])
    antes = Document(BytesIO(ultima['contenido']))
    despues = Document(BytesIO(resultado))
    if _textos(antes) != _textos(despues):
        raise ValueError('El resaltado final alteró el contenido textual')
    verdes = sum(x['color'] == 'verde' for x in evaluacion['contribuciones'])
    amarillas = sum(x['color'] == 'amarillo' for x in evaluacion['contribuciones'])
    return resultado, {'verdes': verdes, 'amarillas': amarillas,
                       'total': verdes + amarillas,
                       'version_evaluada': ultima['numero'],
                       'sha256_word_final': sha256(resultado).hexdigest()}


def validar_cierre_auditoria(ultima, evaluacion, word_final):
    if evaluacion is None or word_final is None:
        return
    esperado, _ = cerrar_autoria(ultima, evaluacion)
    # python-docx puede variar metadatos ZIP al serializar el mismo DOCX.
    # Se cotejan todos los componentes OOXML, no los encabezados del contenedor.
    with ZipFile(BytesIO(esperado)) as a, ZipFile(BytesIO(word_final)) as b:
        if set(a.namelist()) != set(b.namelist()) or any(a.read(n) != b.read(n) for n in a.namelist()):
            raise ValueError('El Word final no coincide con la evaluación documentada')
