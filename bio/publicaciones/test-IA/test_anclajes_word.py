"""Pruebas locales sin consumo de API."""
import io
import json
import unittest
from docx import Document
from versiones import crear_version_original, crear_version_astra
from anclajes_word import obtener_bloques


def documento():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run('La investigación ')
    p.add_run('presenta un ')
    p.add_run('resultado importante.')
    doc.add_paragraph('La investigación presenta un resultado importante.')
    tabla = doc.add_table(rows=1, cols=1)
    tabla.cell(0, 0).text = 'Texto de tabla'
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestAnclajesWord(unittest.TestCase):
    def setUp(self):
        self.base = crear_version_original('original.docx', documento())

    def cambio(self, original, nuevo, ubicacion):
        return {'original': original, 'nuevo': nuevo,
                'motivo': 'Prueba local', 'ubicacion': ubicacion}

    def test_bloques_tienen_identificadores_estables(self):
        bloques = obtener_bloques(self.base['contenido'])
        self.assertEqual([x['id'] for x in bloques], ['P0001', 'P0002', 'T0001R0001C0001P0001'])

    def test_cambio_entre_runs_sin_alterar_otro_parrafo(self):
        nueva = crear_version_astra(self.base, [self.cambio('investigación presenta', 'evaluación describe', 'P0001')],
                                   {'categoria': 'Corrección'})
        doc = Document(io.BytesIO(nueva['contenido']))
        self.assertEqual(doc.paragraphs[0].text, 'La evaluación describe un resultado importante.')
        self.assertEqual(doc.paragraphs[1].text, 'La investigación presenta un resultado importante.')
        self.assertEqual(nueva['base'], 0)

    def test_dos_cambios_mismo_run(self):
        nueva = crear_version_astra(self.base, [self.cambio('La', 'Esta', 'P0001'),
            self.cambio('investigación', 'investigación académica', 'P0001')], {'categoria':'Corrección'})
        doc = Document(io.BytesIO(nueva['contenido']))
        self.assertEqual(doc.paragraphs[0].text, 'Esta investigación académica presenta un resultado importante.')

    def test_traduccion_parrafo_completo(self):
        nueva = crear_version_astra(self.base, [self.cambio('La investigación presenta un resultado importante.',
            'The research presents an important finding.', 'P0001')], {'categoria': 'Traducción'})
        doc = Document(io.BytesIO(nueva['contenido']))
        self.assertEqual(doc.paragraphs[0].text, 'The research presents an important finding.')
        self.assertTrue(all(r.font.highlight_color is None for r in doc.paragraphs[0].runs))

    def test_todas_las_categorias_generan_word_con_cambios_validos(self):
        categorias = ('Traducción', 'Corrección', 'Revisión bibliográfica',
                      'Datos y resultados', 'Argumentación', 'Coherencia integral')
        for categoria in categorias:
            with self.subTest(categoria=categoria):
                nueva = crear_version_astra(self.base, [
                    self.cambio('resultado importante', 'hallazgo relevante', 'P0001')
                ], {'categoria': categoria})
                doc = Document(io.BytesIO(nueva['contenido']))
                self.assertEqual(doc.paragraphs[0].text,
                                 'La investigación presenta un hallazgo relevante.')
                self.assertEqual(nueva['numero'], 1)
                self.assertEqual(nueva['instrucciones']['categoria'], categoria)

    def test_prompt_aplica_cambios_a_todas_las_categorias(self):
        from conexion_astra import preparar_envio_revision
        from iteraciones import crear_solicitud
        for categoria in ('Traducción', 'Corrección', 'Revisión bibliográfica',
                          'Datos y resultados', 'Argumentación', 'Coherencia integral'):
            with self.subTest(categoria=categoria):
                solicitud = crear_solicitud(self.base, categoria, 'Revisar',
                                            'Modificar el texto señalado',
                                            'Manuscrito completo', '')
                paquete = json.loads(preparar_envio_revision(self.base, solicitud))
                self.assertEqual(paquete['categoria'], categoria)
                self.assertIn('cambios', paquete['formato_respuesta'])
                self.assertIn('para cualquier categoría', paquete['regla'].lower())

    def test_tabla(self):
        nueva = crear_version_astra(self.base, [self.cambio('Texto de tabla', 'Table text', 'T0001R0001C0001P0001')],
                                   {'categoria': 'Traducción'})
        doc = Document(io.BytesIO(nueva['contenido']))
        self.assertEqual(doc.tables[0].cell(0, 0).text, 'Table text')

    def test_ambiguo_no_modifica_original(self):
        with self.assertRaises(ValueError):
            crear_version_astra(self.base, [self.cambio('investigación', 'research', 'párrafo')], {})
        self.assertEqual(Document(io.BytesIO(self.base['contenido'])).paragraphs[0].text,
                         'La investigación presenta un resultado importante.')

    def test_solapamiento_rechazado(self):
        with self.assertRaises(ValueError):
            crear_version_astra(self.base, [self.cambio('investigación presenta', 'X', 'P0001'),
              self.cambio('presenta un', 'Y', 'P0001')], {})


    def test_id_inexistente_no_utiliza_busqueda_global(self):
        with self.assertRaisesRegex(ValueError, 'ID de bloque'):
            crear_version_astra(self.base, [self.cambio('resultado importante', 'hallazgo', 'P9999')],
                                {'categoria': 'Argumentación'})

    def test_formato_de_tramos_no_modificados(self):
        doc = Document()
        p = doc.add_paragraph()
        r1 = p.add_run('Prefijo '); r1.bold = True
        r2 = p.add_run('texto a revisar'); r2.italic = True
        r3 = p.add_run(' sufijo'); r3.underline = True
        buf = io.BytesIO(); doc.save(buf)
        base = crear_version_original('formato.docx', buf.getvalue())
        nueva = crear_version_astra(base, [self.cambio('texto a revisar', 'texto nuevo', 'P0001')],
                                    {'categoria': 'Argumentación'})
        parrafo = Document(io.BytesIO(nueva['contenido'])).paragraphs[0]
        self.assertEqual(parrafo.text, 'Prefijo texto nuevo sufijo')
        self.assertTrue(next(r for r in parrafo.runs if r.text == 'Prefijo ').bold)
        self.assertTrue(next(r for r in parrafo.runs if r.text == ' sufijo').underline)
        self.assertTrue(next(r for r in parrafo.runs if r.text == 'texto nuevo').italic)

    def test_dos_cambios_mismo_run_con_sufijo(self):
        doc = Document(); doc.add_paragraph('uno dos tres cuatro')
        buf = io.BytesIO(); doc.save(buf)
        base = crear_version_original('simple.docx', buf.getvalue())
        nueva = crear_version_astra(base, [self.cambio('dos', 'DOS', 'P0001'),
                                          self.cambio('cuatro', 'CUATRO', 'P0001')],
                                    {'categoria': 'Argumentación'})
        self.assertEqual(Document(io.BytesIO(nueva['contenido'])).paragraphs[0].text,
                         'uno DOS tres CUATRO')

    def test_rechazo_xml_complejo_no_altera_original(self):
        doc = Document(); p = doc.add_paragraph('texto')
        p.add_run().add_break()
        buf = io.BytesIO(); doc.save(buf)
        base = crear_version_original('complejo.docx', buf.getvalue())
        with self.assertRaises(ValueError):
            crear_version_astra(base, [self.cambio('texto', 'otro', 'P0001')],
                                {'categoria': 'Argumentación'})
        self.assertEqual(Document(io.BytesIO(base['contenido'])).paragraphs[0].text, 'texto\n')

    def test_payload_no_duplica_texto(self):
        from conexion_astra import preparar_envio_revision
        from iteraciones import crear_solicitud
        solicitud = crear_solicitud(self.base, 'Argumentación', 'Revisar', 'Mejorar argumento',
                                    'Manuscrito completo', '')
        paquete = json.loads(preparar_envio_revision(self.base, solicitud))
        self.assertNotIn('bloques_word', paquete)
        self.assertEqual(paquete['contenido_word'].count('La investigación presenta un resultado importante.'), 2)
        self.assertIn('[P0001]', paquete['contenido_word'])

    def test_cambios_diferentes_bloques_preservan_otros(self):
        nueva = crear_version_astra(self.base, [
            self.cambio('resultado importante', 'hallazgo significativo', 'P0001'),
            self.cambio('Texto de tabla', 'Table content', 'T0001R0001C0001P0001')],
            {'categoria': 'Argumentación'})
        doc = Document(io.BytesIO(nueva['contenido']))
        self.assertEqual(doc.paragraphs[1].text, 'La investigación presenta un resultado importante.')
        self.assertEqual(doc.tables[0].cell(0, 0).text, 'Table content')


if __name__ == '__main__':
    unittest.main()
