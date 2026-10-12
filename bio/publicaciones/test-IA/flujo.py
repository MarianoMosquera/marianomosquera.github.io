"""Operaciones locales comprobables. Ninguna llama a servicios externos."""
from io import BytesIO
from difflib import SequenceMatcher
from hashlib import sha256
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from versiones import crear_version, validar_docx
from iteraciones import crear_solicitud
from anotaciones import crear_documento_final


def solicitar_iteracion(base, categoria, general, especifica, alcance='Manuscrito completo', ubicacion=''):
    """Registra instrucciones nuevas vinculadas a la versión vigente."""
    evento = crear_solicitud(base, categoria, general, especifica, alcance, ubicacion)
    evento['tipo'] = 'Nueva iteración del AUTOR'
    return evento


def comparar_offline(base_bytes, nuevo_bytes):
    """Comparación informativa de párrafos; NO es prueba de autoría."""
    validar_docx(base_bytes)
    validar_docx(nuevo_bytes)
    a = [p.text for p in Document(BytesIO(base_bytes)).paragraphs]
    b = [p.text for p in Document(BytesIO(nuevo_bytes)).paragraphs]
    cambios = []
    for op, i, j, k, l in SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if op != 'equal':
            cambios.append({'operacion': op, 'parrafos_anteriores': [i, j],
                            'parrafos_nuevos': [k, l],
                            'texto_anterior': a[i:j], 'texto_nuevo': b[k:l]})
    return cambios


def crear_traduccion_completa(base, docx_traducido, instrucciones):
    """Una traducción completa se registra como versión de Astra, sin verde masivo."""
    return crear_version(base, docx_traducido, 'astra', instrucciones=instrucciones)


def agregar_evaluacion_coherencia(base, informe, instrucciones):
    """Informe evaluativo al final, en verde, sin tocar el cuerpo del texto."""
    if not informe or not informe.strip():
        raise ValueError('Falta el informe de coherencia integral')
    doc = Document(BytesIO(base['contenido']))
    doc.add_page_break()
    titulo = doc.add_paragraph('Evaluación de coherencia integral', style='Heading 1')
    for run in titulo.runs:
        run.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
    for linea in informe.strip().splitlines():
        parrafo = doc.add_paragraph()
        parrafo.add_run(linea).font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
    buf = BytesIO()
    doc.save(buf)
    return crear_version(base, buf.getvalue(), 'astra', instrucciones=instrucciones)


def validar_evaluacion_final(evaluacion, contenido):
    """Exige evaluación externa explícita y anclajes verificables."""
    if not isinstance(evaluacion, dict) or not isinstance(evaluacion.get('contribuciones'), list):
        raise ValueError('Falta una evaluación estructurada de Astra')
    for item in evaluacion['contribuciones']:
        if item.get('color') not in ('verde', 'amarillo'):
            raise ValueError('Color inválido')
        if not item.get('texto') or not item.get('justificacion'):
            raise ValueError('Cada contribución necesita texto y justificación')
        if item['color'] == 'amarillo':
            cuartil = item.get('cuartil')
            if cuartil is not None:
                if cuartil not in ('Q1', 'Q2', 'Q3', 'Q4'):
                    raise ValueError('Cuartil amarillo inválido')
            else:
                # Compatibilidad con expedientes anteriores; nuevas evaluaciones usan cuartiles.
                porcentaje = item.get('porcentaje_estimado')
                if isinstance(porcentaje, bool) or not isinstance(porcentaje, (int, float)) or not 0 <= porcentaje <= 100:
                    raise ValueError('El amarillo requiere cuartil Q1, Q2, Q3 o Q4')
    # Verificar anclajes sin inventar autoría. El Word se crea únicamente si todos son únicos.
    word = crear_documento_final(contenido, evaluacion['contribuciones'])
    return word
