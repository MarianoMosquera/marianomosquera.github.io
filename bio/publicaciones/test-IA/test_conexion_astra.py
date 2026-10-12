import io
import json
import unittest
from unittest.mock import Mock
from docx import Document
from versiones import crear_version_original
from iteraciones import crear_solicitud
from conexion_astra import (extraer_word, preparar_envio_revision, preparar_envio_final,
    ejecutar_api, ConexionErrorSegura, interpretar_json)
from integracion import procesar_revision


def word():
    d = Document()
    d.add_paragraph('El estudio presenta un resultado.')
    t = d.add_table(rows=1, cols=2)
    t.cell(0,0).text='Variable'
    t.cell(0,1).text='Valor'
    b=io.BytesIO(); d.save(b); return b.getvalue()

class TestConexion(unittest.TestCase):
    def setUp(self):
        self.base = crear_version_original('prueba.docx', word())
        self.solicitud = crear_solicitud(self.base, 'Corrección', 'Corregí estilo',
                                         'Mejorar redacción', 'Manuscrito completo', '')

    def test_incluye_texto_y_tablas_reales(self):
        p=json.loads(preparar_envio_revision(self.base,self.solicitud))
        self.assertIn('El estudio presenta un resultado.',p['contenido_word'])
        self.assertIn('[T0001R0001C0001P0001] Variable',p['contenido_word'])
        self.assertIn('[T0001R0001C0002P0001] Valor',p['contenido_word'])
        self.assertNotIn('bloques_word', p)
        self.assertNotIn('OPENAI_API_KEY',json.dumps(p))

    def test_version_inconsistente_no_se_envia(self):
        self.solicitud['base_sha256']='falso'
        with self.assertRaises(ValueError): preparar_envio_revision(self.base,self.solicitud)

    def test_clave_ausente_no_hace_llamada(self):
        client=Mock()
        with self.assertRaises(ConexionErrorSegura):
            ejecutar_api('{}',api_key='',client=client)
        client.responses.create.assert_not_called()

    def test_store_false_y_respuesta(self):
        client=Mock()
        client.responses.create.return_value.output_text='{"cambios": [], "informe": "Sin cambios"}'
        r=ejecutar_api('{}',api_key='clave-ficticia',client=client)
        self.assertIn('Sin cambios',r)
        self.assertFalse(client.responses.create.call_args.kwargs['store'])

    def test_revision_y_trazabilidad(self):
        client=Mock()
        respuesta={'cambios':[{'original':'resultado','nuevo':'hallazgo',
                    'motivo':'Precisión','ubicacion':'P0001'}],'informe':'Corrección puntual'}
        client.responses.create.return_value.output_text=json.dumps(respuesta)
        e,n,aviso=procesar_revision(self.base,self.solicitud,api_key='ficticia',cliente=client)
        self.assertIsNotNone(n)
        self.assertEqual(e['version_resultante'],1)
        self.assertEqual(e['ejecucion']['respuesta_completa'],json.dumps(respuesta))

    def test_anclaje_ambiguo_conserva_respuesta_sin_version(self):
        client=Mock()
        client.responses.create.return_value.output_text=json.dumps({'cambios':[
            {'original':'no-existe','nuevo':'otro','motivo':'x','ubicacion':'p'}], 'informe':''})
        e,n,aviso=procesar_revision(self.base,self.solicitud,api_key='ficticia',cliente=client)
        self.assertIsNone(n)
        self.assertIn('rechazados',e['nota_aplicacion'])
        self.assertIsNone(e['version_resultante'])
        from exportacion import crear_zip_auditoria
        from io import BytesIO
        from zipfile import ZipFile
        archivo = crear_zip_auditoria(self.base, [], [e])
        with ZipFile(BytesIO(archivo)) as z:
            self.assertIsNone(z.testzip())
            self.assertIn('documentos/original.docx', z.namelist())

    def test_error_proveedor_sin_detalles(self):
        client=Mock()
        client.responses.create.side_effect=RuntimeError('API KEY SECRETA 123')
        with self.assertRaises(ConexionErrorSegura) as cm:
            ejecutar_api('{}',api_key='ficticia',client=client)
        self.assertNotIn('SECRETA',str(cm.exception))

    def test_final_contiene_todas_las_versiones(self):
        p=json.loads(preparar_envio_final(self.base,[],[]))
        self.assertEqual([x['numero'] for x in p['versiones']],[0])
        self.assertIn('El estudio presenta',p['versiones'][0]['texto'])

if __name__=='__main__': unittest.main()
